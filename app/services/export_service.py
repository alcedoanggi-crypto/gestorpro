"""Exportacion de reportes a Excel (.xlsx) y PDF."""
from io import BytesIO
from datetime import datetime, date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.colors import Color, HexColor
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from ..models.enums import TaskStatus

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(color="FFFFFF", bold=True)

GANTT_STATUS_COLORS = {
    TaskStatus.TODO: HexColor("#94a3b8"),
    TaskStatus.IN_PROGRESS: HexColor("#2563eb"),
    TaskStatus.REVIEW: HexColor("#f97316"),
    TaskStatus.DONE: HexColor("#10b981"),
}


def _fmt_date(d):
    return d.strftime("%d/%m/%Y") if d else "-"


# --------------------------------------------------------------------------- #
# EXCEL
# --------------------------------------------------------------------------- #
def projects_to_excel(projects):
    wb = Workbook()
    ws = wb.active
    ws.title = "Proyectos"
    headers = ["ID", "Nombre", "Estado", "Prioridad", "Gerente", "Equipo",
               "Inicio", "Fin", "% Avance", "Tareas", "Abiertas", "Vencidas"]
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")

    for p in projects:
        ws.append([
            p.id, p.name, p.status_label, p.priority_label,
            p.manager.name if p.manager else "-",
            p.team.name if p.team else "-",
            _fmt_date(p.start_date), _fmt_date(p.end_date),
            p.progress, p.task_count, p.open_task_count, p.overdue_task_count,
        ])

    # Hoja de tareas
    ws2 = wb.create_sheet("Tareas")
    h2 = ["ID", "Proyecto", "Tarea", "Estado", "Prioridad", "Responsable",
          "Inicio", "Vence", "% Avance", "Vencida"]
    ws2.append(h2)
    for c in ws2[1]:
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")
    for p in projects:
        for t in p.tasks:
            ws2.append([
                t.id, p.name, t.title, t.status_label, t.priority_label,
                t.assignee.name if t.assignee else "-",
                _fmt_date(t.start_date), _fmt_date(t.due_date),
                t.progress, "Si" if t.is_overdue else "No",
            ])

    for sheet in (ws, ws2):
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 3, 45)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def report_to_pdf(projects, stats):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    elems = [
        Paragraph("Reporte de Proyectos - GestorPro", styles["Title"]),
        Paragraph(datetime.now().strftime("Generado el %d/%m/%Y %H:%M"), styles["Normal"]),
        Spacer(1, 0.6 * cm),
    ]

    resumen = [
        ["Proyectos", "Activos", "Retrasados", "Tareas pendientes", "Tareas vencidas"],
        [stats["total_projects"], stats["active_projects"], stats["delayed_projects"],
         stats["pending_tasks"], stats["overdue_tasks"]],
    ]
    t = Table(resumen, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elems += [t, Spacer(1, 0.8 * cm),
              Paragraph("Detalle de proyectos", styles["Heading2"])]

    data = [["Nombre", "Estado", "Prioridad", "Gerente", "Inicio", "Fin", "% Avance", "Tareas"]]
    for p in projects:
        data.append([
            Paragraph(p.name, styles["BodyText"]), p.status_label, p.priority_label,
            p.manager.name if p.manager else "-",
            _fmt_date(p.start_date), _fmt_date(p.end_date),
            f"{p.progress}%", p.task_count,
        ])
    dt = Table(data, repeatRows=1, colWidths=[6 * cm, 2.6 * cm, 2.2 * cm, 4 * cm,
                                              2.6 * cm, 2.6 * cm, 2 * cm, 1.8 * cm])
    dt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
    ]))
    elems.append(dt)

    doc.build(elems)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------- #
# PDF - Cronograma (Gantt) por proyecto
# --------------------------------------------------------------------------- #
def _task_range(t):
    start = t.start_date or t.due_date or date.today()
    end = t.due_date or t.start_date or date.today()
    return start, (end if end >= start else start)


def _darken(c, f=0.55):
    return Color(c.red * f, c.green * f, c.blue * f)


def _gantt_legend():
    d = Drawing(700, 14)
    x = 0
    for st in TaskStatus.ALL:
        label = TaskStatus.LABELS[st]
        d.add(Rect(x, 2, 10, 10, fillColor=GANTT_STATUS_COLORS[st], strokeColor=None))
        d.add(String(x + 14, 3, label, fontSize=7, fillColor=HexColor("#334155")))
        x += 22 + 4.6 * len(label) + 14
    return d


def _gantt_drawing(tasks, width, dmin, dmax, row_h=19, label_w=175):
    total_days = max((dmax - dmin).days, 1)
    chart_w = width - label_w
    height = row_h * len(tasks) + 26
    top = height - 22
    d = Drawing(width, height)

    def x_of(dt):
        return label_w + (dt - dmin).days / total_days * chart_w

    # --- Ejes de tiempo (semanal si el rango es corto, mensual si es largo) ---
    ticks = []
    if total_days <= 70:
        cur = dmin - timedelta(days=dmin.weekday())
        while cur <= dmax:
            if cur >= dmin:
                ticks.append((cur, cur.strftime("%d/%m")))
            cur += timedelta(days=7)
    else:
        cur = date(dmin.year, dmin.month, 1)
        while cur <= dmax:
            if cur >= dmin:
                ticks.append((cur, cur.strftime("%b %Y")))
            year, month = divmod(cur.month, 12)
            cur = date(cur.year + year, month + 1, 1)

    for tk, text in ticks:
        xx = x_of(tk)
        d.add(Line(xx, 0, xx, top, strokeColor=HexColor("#e2e8f0"), strokeWidth=0.5))
        d.add(String(xx + 2, top + 6, text, fontSize=6.5, fillColor=HexColor("#64748b")))

    today = date.today()
    if dmin <= today <= dmax:
        xx = x_of(today)
        d.add(Line(xx, 0, xx, top + 2, strokeColor=HexColor("#ef4444"),
                   strokeWidth=0.9, strokeDashArray=[2, 2]))
        d.add(String(xx + 2, 1, "hoy", fontSize=6, fillColor=HexColor("#ef4444")))

    d.add(Line(label_w, 0, label_w, top, strokeColor=HexColor("#cbd5e1"), strokeWidth=0.6))

    for i, t in enumerate(tasks):
        y = top - (i + 1) * row_h
        start, end = _task_range(t)
        x1, x2 = x_of(start), x_of(end)
        bar_w = max(x2 - x1, 2.5)
        color = GANTT_STATUS_COLORS.get(t.status, HexColor("#94a3b8"))

        name = t.title if len(t.title) <= 33 else t.title[:32] + "…"
        d.add(String(4, y + 6, name, fontSize=7, fillColor=HexColor("#1e293b")))
        d.add(Rect(label_w, y + 3, chart_w, row_h - 7,
                   fillColor=HexColor("#f1f5f9"), strokeColor=None))
        d.add(Rect(x1, y + 3, bar_w, row_h - 7, rx=2, ry=2,
                   fillColor=color, strokeColor=None))
        pct = max(0, min(t.progress or 0, 100))
        if pct:
            d.add(Rect(x1, y + 3, bar_w * pct / 100.0, row_h - 7, rx=2, ry=2,
                       fillColor=_darken(color), strokeColor=None))
        d.add(String(min(x1 + bar_w + 3, width - 22), y + 6, f"{pct}%",
                     fontSize=6, fillColor=HexColor("#64748b")))
    return d


def project_gantt_to_pdf(project):
    buffer = BytesIO()
    page_w = landscape(A4)[0]
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.3 * cm, bottomMargin=1.3 * cm,
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    usable_w = page_w - 2.4 * cm

    elems = [
        Paragraph(f"Cronograma (Gantt) - {project.name}", styles["Title"]),
        Paragraph(datetime.now().strftime("Generado el %d/%m/%Y %H:%M"), styles["Normal"]),
    ]

    tasks = sorted(
        project.tasks,
        key=lambda t: (_task_range(t)[0], t.id),
    )
    if not tasks:
        elems.append(Spacer(1, 0.6 * cm))
        elems.append(Paragraph("Este proyecto todavia no tiene tareas.", styles["Normal"]))
        doc.build(elems)
        buffer.seek(0)
        return buffer

    dmin = min(_task_range(t)[0] for t in tasks)
    dmax = max(_task_range(t)[1] for t in tasks)
    if dmin == dmax:
        dmax = dmin + timedelta(days=7)
    pad = max(1, round((dmax - dmin).days * 0.04))
    dmin -= timedelta(days=pad)
    dmax += timedelta(days=pad)

    elems += [Spacer(1, 0.25 * cm), _gantt_legend(), Spacer(1, 0.2 * cm)]

    per_page = 18
    for i in range(0, len(tasks), per_page):
        elems.append(_gantt_drawing(tasks[i:i + per_page], usable_w, dmin, dmax))
        if i + per_page < len(tasks):
            elems.append(PageBreak())

    doc.build(elems)
    buffer.seek(0)
    return buffer
